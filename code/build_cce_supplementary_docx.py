"""Build a journal-ready Supplementary Information DOCX from the frozen SI draft."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RESULTS = ROOT / "10_TEA_LCA_Python" / "results" / "pycharm_publication_enhancement"
SOURCE = RESULTS / "PSA_QAE_CCE_SUPPLEMENTARY_INFORMATION_DRAFT.md"
OUTPUT = RESULTS / "PSA_QAE_CCE_SUPPLEMENTARY_INFORMATION.docx"
TITLE = (
    "Resource-aware quantum amplitude estimation for chance-constrained pressure swing "
    "adsorption: classical baselines, compiled oracle costs, and a hardware boundary"
)
FULL_TABLES = (
    (
        "S1",
        RESULTS / "PSA_QAE_CCE_RESOURCE_TABLE3.csv",
        "Complete 31-tier quantum-resource evidence ladder",
        [1.05, 1.35, 0.50, 0.85, 1.20, 1.35, 1.25, 2.20],
    ),
    (
        "S2",
        RESULTS / "PSA_QAE_CCE_TABLE4_EVIDENCE_HIERARCHY.csv",
        "Complete 32-tier process, validation, and hardware evidence hierarchy",
        [1.45, 1.6, 1.35, 1.65, 1.55, 2.15],
    ),
)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def _clean(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_margins(cell, top=10, start=55, bottom=10, end=55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_border(cell, *, top=False, bottom=False) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        enabled = (edge == "top" and top) or (edge == "bottom" and bottom)
        el.set(qn("w:val"), "single" if enabled else "nil")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "777777")


def _font_run(run, size=6.8, bold=False, italic=False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _add_full_table(doc: Document, label: str, csv_path: Path, title: str, widths: list[float]) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)

    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f"Supplementary Table {label}. {title}")
    _font_run(run, size=10, bold=True)

    with csv_path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if len(widths) != len(rows[0]):
        raise ValueError(f"{label} width profile mismatch")

    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    _repeat_header(table.rows[0])
    _keep_row_together(table.rows[0])

    for i, text in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_border(cell, top=True, bottom=True)
        _set_cell_margins(cell, top=18, bottom=18)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        _font_run(para.add_run(text), size=6.9, bold=True)

    for row_index, row in enumerate(rows[1:], start=1):
        table_row = table.add_row()
        _keep_row_together(table_row)
        cells = table_row.cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_border(cell, bottom=row_index == len(rows) - 1)
            _set_cell_margins(cell)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            _font_run(para.add_run(text), size=6.45)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.50)
    section.bottom_margin = Inches(0.50)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.35)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color in (("Heading 1", 14, "17365D"), ("Heading 2", 11.5, "17365D")):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.space_before = Pt(6.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True

    if "SI Bullet" not in [s.name for s in doc.styles]:
        bullet = doc.styles.add_style("SI Bullet", WD_STYLE_TYPE.PARAGRAPH)
        bullet.base_style = normal
        bullet.paragraph_format.left_indent = Inches(0.22)
        bullet.paragraph_format.first_line_indent = Inches(-0.18)
        bullet.paragraph_format.space_after = Pt(1.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Supplementary Information · PSA–QAE resource audit")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(90, 90, 90)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("SI-")
    fr.font.name = "Arial"
    fr.font.size = Pt(8)
    _field(footer, "PAGE")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Supplementary Information")
    r.font.name = "Arial"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = RGBColor(23, 54, 93)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(TITLE)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.5)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Ugwiyeon Lee*, Seong-Yun Bae, Hakrae Kim, Jungsu Lee")
    r.bold = True
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run(
        "Korea Gas Safety Corporation, 1390 Wonjung-ro, Maengdong-myeon, "
        "Eumseong-gun, Chungcheongbuk-do 27738, Republic of Korea"
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.add_run("* Corresponding author: Ugwiyeon Lee (lyullee@kgs.or.kr)").italic = True

    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "EAF2F8")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Claim boundary: this supplement reports conditional process-model, estimator, "
        "circuit-resource, and hardware evidence. It does not establish quantum advantage "
        "or industrial hydrogen validation."
    )
    r.bold = True
    r.font.size = Pt(9.5)

    # Skip the source's generic title; retain all numbered sections and evidence gates.
    paragraph_buffer: list[str] = []

    def flush() -> None:
        if not paragraph_buffer:
            return
        text = _clean(" ".join(paragraph_buffer))
        paragraph_buffer.clear()
        if text:
            doc.add_paragraph(text)

    for line in lines[1:]:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("## "):
            flush()
            doc.add_heading(_clean(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            flush()
            doc.add_heading(_clean(stripped[4:]), level=2)
        elif stripped.startswith("- "):
            flush()
            p = doc.add_paragraph(style="SI Bullet")
            p.add_run("• ")
            p.add_run(_clean(stripped[2:]))
        else:
            paragraph_buffer.append(stripped)
    flush()

    for label, csv_path, title, widths in FULL_TABLES:
        _add_full_table(doc, label, csv_path, title, widths)

    # Word reserves a terminal paragraph after a table.  Make that paragraph
    # explicit and one point high so the final landscape table does not push an
    # otherwise empty page into the rendered supplement.
    terminal = doc.add_paragraph()
    terminal.paragraph_format.space_before = Pt(0)
    terminal.paragraph_format.space_after = Pt(0)
    terminal.paragraph_format.line_spacing = Pt(1)
    terminal_run = terminal.add_run("\u200b")
    terminal_run.font.size = Pt(1)

    props = doc.core_properties
    props.title = "Supplementary Information — resource-aware PSA QAE"
    props.subject = "Computers & Chemical Engineering submission supplement"
    props.author = "Ugwiyeon Lee; Seong-Yun Bae; Hakrae Kim; Jungsu Lee"
    props.keywords = "pressure swing adsorption; quantum amplitude estimation; supplementary information"
    doc.save(OUTPUT)
    print(f"wrote {OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")


if __name__ == "__main__":
    build()

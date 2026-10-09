"""Build the separate CCE title page from validated author metadata."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RESULTS = ROOT / "10_TEA_LCA_Python" / "results" / "pycharm_publication_enhancement"
SOURCE = RESULTS / "PSA_QAE_CCE_CORE_MANUSCRIPT_DRAFT.md"
AUTHOR_FILE = RESULTS / "PSA_QAE_CCE_AUTHOR_METADATA.json"
OUTPUT = RESULTS / "PSA_QAE_CCE_TITLE_PAGE.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_font(run, *, size: float = 10.5, bold: bool = False, italic: bool = False, color="000000") -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text: str, **style):
    run = paragraph.add_run(text)
    set_font(run, **style)
    return run


def main() -> None:
    source_text = SOURCE.read_text(encoding="utf-8")
    title_match = re.search(r"^#\s+(.+)$", source_text, re.M)
    if not title_match:
        raise RuntimeError("manuscript title not found")
    title = title_match.group(1).strip()
    data = json.loads(AUTHOR_FILE.read_text(encoding="utf-8"))
    authors = sorted(data["authors"], key=lambda item: item["order"])
    corresponding = next(author for author in authors if author["corresponding_author"])

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    add_text(p, title, size=16, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    for index, author in enumerate(authors):
        if index:
            add_text(p, ", ", size=11.5)
        add_text(p, author["name_en"], size=11.5, bold=True)
        if author["corresponding_author"]:
            marker = add_text(p, "*", size=11.5, bold=True)
            marker.font.superscript = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    add_text(
        p,
        "Korea Gas Safety Corporation, 1390 Wonjung-ro, Maengdong-myeon, "
        "Eumseong-gun, Chungcheongbuk-do 27738, Republic of Korea",
        size=10.5,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    add_text(p, "* Corresponding author: ", size=10.5, bold=True)
    add_text(p, f"{corresponding['name_en']} ({corresponding['email']})", size=10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    add_text(p, "Author identifiers", size=12.5, bold=True)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.55, 2.25, 3.55]
    headers = ["Author", "ORCID", "Email"]
    for col, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[col]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "1F4E78")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = add_text(p, header, size=9.5, bold=True, color="FFFFFF")
    for row_index, author in enumerate(authors, start=1):
        cells = table.add_row().cells
        values = [author["name_en"], author["orcid"], author["email"]]
        for col, (value, width) in enumerate(zip(values, widths)):
            cells[col].width = Inches(width)
            cells[col].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cells[col], "EAF2F8")
            p = cells[col].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, size=9.25)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    add_text(p, "CRediT authorship contribution statement", size=12.5, bold=True)

    role_parts = []
    for author in authors:
        roles = ", ".join(author["credit_roles"])
        role_parts.append(f"{author['name_en']}: {roles}.")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.05
    add_text(p, " ".join(role_parts), size=10.25)

    doc.core_properties.title = f"Title page — {title}"
    doc.core_properties.subject = "Computers & Chemical Engineering submission title page"
    doc.core_properties.author = corresponding["name_en"]
    doc.core_properties.keywords = "pressure swing adsorption; quantum amplitude estimation"
    doc.save(OUTPUT)

    all_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    all_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    for author in authors:
        assert author["name_en"] in all_text
        assert author["orcid"] in all_text
        assert author["email"] in all_text
    assert corresponding["name_en"] in all_text and corresponding["email"] in all_text
    print(f"wrote {OUTPUT}")
    print(f"authors={len(authors)} tables={len(doc.tables)} sections={len(doc.sections)}")


if __name__ == "__main__":
    main()

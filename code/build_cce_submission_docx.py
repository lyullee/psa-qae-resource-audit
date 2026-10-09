from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RESULTS = ROOT / "10_TEA_LCA_Python" / "results" / "pycharm_publication_enhancement"
SOURCE = RESULTS / "PSA_QAE_CCE_CORE_MANUSCRIPT_DRAFT.md"
BIB = RESULTS / "PSA_QAE_SCI_REFERENCES.bib"
CAPTIONS = RESULTS / "PSA_QAE_CCE_MAIN_FIGURE_CAPTIONS.md"
OUTPUT = RESULTS / "PSA_QAE_CCE_SUBMISSION_MANUSCRIPT.docx"

FIGURES = {
    1: RESULTS / "PSA_QAE_CCE_WORKFLOW_FIGURE.png",
    2: RESULTS / "PSA_QAE_CONTINUOUS_POLICY_SELECTION_FIGURE.png",
    3: RESULTS / "PSA_QAE_CCE_PRIMARY_RESOURCE_BOUNDARY_FIGURE.png",
    4: RESULTS / "PSA_QAE_CCE_RESOURCE_LADDER_FIGURE.png",
    5: RESULTS / "FIG_QPU_MLAE_HARDWARE_BOUNDARY.png",
    6: RESULTS / "PSA_QAE_CCE_VALIDITY_BOUNDARY_FIGURE.png",
}

TABLES = {
    1: (RESULTS / "PSA_QAE_CCE_TABLE1_FROZEN_PROTOCOL.csv", "Frozen primary protocol and claim boundary"),
    2: (RESULTS / "PSA_QAE_CCE_TABLE2_MATCHED_ESTIMATORS.csv", "Matched estimator performance over 200 replications"),
    3: (RESULTS / "PSA_QAE_CCE_RESOURCE_TABLE3.csv", "Condensed quantum-resource evidence ladder"),
    4: (RESULTS / "PSA_QAE_CCE_TABLE4_EVIDENCE_HIERARCHY.csv", "Condensed process, validation, and hardware evidence hierarchy"),
}

MAIN_TABLE_SELECTION = {
    3: {
        "finite-table phase oracle",
        "fixed-point response-event oracle",
        "dense empirical-copula loader (2 bit)",
        "primary policy stream-to-event clean oracle",
        "gas-inventory-generated product-cell LDF/valve/balance paths",
        "coherent gas-inventory-to-dynamic-Langmuir/transport/LDF/balance paths",
        "pressure-derived saturated/unsaturated proportional tank-allocation interface",
        "real-QPU three-qubit event",
    },
    4: {
        "primary spatial/tank response model",
        "primary physical CSS horizon and register accounting",
        "primary policy stream-to-event clean oracle",
        "coherent gas-inventory-to-dynamic-Langmuir/transport/LDF/balance paths",
        "pressure-derived saturated/unsaturated proportional tank-allocation interface",
        "Thelen VA-10 matched H2 pair",
        "public toPSAil/Yang H2 model and experiment",
        "real-QPU finite event",
    },
}


def set_run_font(run, name="Times New Roman", size=10.5, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + side
        node = tcMar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, *, top=False, bottom=False, color="777777", size="8"):
    """Apply CCE-style horizontal rules while explicitly suppressing vertical rules."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        enabled = (edge == "top" and top) or (edge == "bottom" and bottom)
        el.set(qn("w:val"), "single" if enabled else "nil")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=9)


def set_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)
    # Continuous line numbers every five lines aid review.
    sectPr = section._sectPr
    ln = sectPr.find(qn("w:lnNumType"))
    if ln is None:
        ln = OxmlElement("w:lnNumType")
        sectPr.append(ln)
    ln.set(qn("w:countBy"), "5")
    ln.set(qn("w:restart"), "continuous")
    # Header and footer.
    header = section.header
    p = header.paragraphs[0]
    p.text = "Resource-aware QAE for pressure swing adsorption"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        set_run_font(run, size=8.5, italic=True, color="555555")
    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()
    add_page_number(p)


def clean_tex(text: str) -> str:
    replacements = {
        r'{\"a}': "ä", r'{\"o}': "ö", r'{\"u}': "ü",
        r'{\'c}': "ć", r"{\'e}": "é", r"{\'a}": "á",
        r"{\v{c}}": "č", r"{\v{s}}": "š", r"{\v{z}}": "ž",
        r"\&": "&", "{": "", "}": "",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.replace("\\", "")


def parse_bib(path: Path) -> dict[str, dict[str, str]]:
    raw = path.read_text(encoding="utf-8")
    entries: dict[str, dict[str, str]] = {}
    pos = 0
    while True:
        match = re.search(r"@(\w+)\s*\{\s*([^,]+),", raw[pos:], re.S)
        if not match:
            break
        start = pos + match.start()
        body_start = pos + match.end()
        depth = 1
        i = body_start
        while i < len(raw) and depth:
            if raw[i] == "{":
                depth += 1
            elif raw[i] == "}":
                depth -= 1
            i += 1
        entry_type = match.group(1).lower()
        key = match.group(2).strip()
        body = raw[body_start:i - 1]
        fields: dict[str, str] = {"ENTRYTYPE": entry_type}
        j = 0
        while j < len(body):
            fm = re.search(r"([A-Za-z][A-Za-z0-9_-]*)\s*=\s*", body[j:])
            if not fm:
                break
            name = fm.group(1).lower()
            k = j + fm.end()
            if k >= len(body):
                break
            if body[k] == "{":
                d = 1
                q = k + 1
                while q < len(body) and d:
                    if body[q] == "{": d += 1
                    elif body[q] == "}": d -= 1
                    q += 1
                value = body[k + 1:q - 1]
                j = q
            elif body[k] == '"':
                q = k + 1
                while q < len(body) and body[q] != '"': q += 1
                value = body[k + 1:q]
                j = q + 1
            else:
                q = body.find(",", k)
                if q < 0: q = len(body)
                value = body[k:q].strip()
                j = q + 1
            fields[name] = clean_tex(value.strip())
        entries[key] = fields
        pos = i
    return entries


def split_authors(entry: dict[str, str]) -> list[str]:
    raw = entry.get("author", entry.get("editor", "Unknown author"))
    return [part.strip() for part in raw.split(" and ") if part.strip()]


def author_surname(author: str) -> str:
    if author.lower() == "others":
        return "et al."
    if "," in author:
        return author.split(",", 1)[0].strip()
    parts = author.split()
    return parts[-1] if parts else "Unknown"


def author_initials(author: str) -> str:
    if "," in author:
        given = author.split(",", 1)[1].strip()
    else:
        given = " ".join(author.split()[:-1])
    tokens = re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]+", given)
    return ".".join(token[0].upper() for token in tokens) + ("." if tokens else "")


def reference_authors(entry: dict[str, str]) -> str:
    formatted: list[str] = []
    for author in split_authors(entry):
        if author.lower() == "others":
            formatted.append("et al.")
        else:
            initials = author_initials(author)
            formatted.append(f"{author_surname(author)}, {initials}" if initials else author_surname(author))
    return ", ".join(formatted)


def citation_label(entry: dict[str, str]) -> str:
    authors = [a for a in split_authors(entry) if a.lower() != "others"]
    surnames = [author_surname(a) for a in authors]
    if len(surnames) == 1:
        name = surnames[0]
    elif len(surnames) == 2:
        name = f"{surnames[0]} and {surnames[1]}"
    else:
        name = f"{surnames[0]} et al."
    return f"{name}, {entry.get('year', 'n.d.')}"


def reference_sort_key(entry: dict[str, str]) -> tuple[str, str, str]:
    first = next((a for a in split_authors(entry) if a.lower() != "others"), "Unknown")
    return (author_surname(first).casefold(), entry.get("year", "n.d."), entry.get("title", "").casefold())


def reference_text(entry: dict[str, str]) -> str:
    authors = reference_authors(entry)
    title = entry.get("title", "Untitled")
    year = entry.get("year", "n.d.")
    kind = entry.get("ENTRYTYPE", "")
    if kind == "article":
        journal = entry.get("journal", "")
        volume = entry.get("volume", "")
        number = entry.get("number", "")
        pages = entry.get("pages", entry.get("page", ""))
        venue = journal
        if volume:
            venue += f" {volume}"
        if number:
            venue += f"({number})"
        if pages:
            venue += f", {pages}"
        core = f"{authors}, {year}. {title}. {venue}."
    elif kind in {"inproceedings", "conference"}:
        core = f"{authors}, {year}. {title}. {entry.get('booktitle', '')}."
    elif kind in {"phdthesis", "mastersthesis", "thesis"}:
        inst = entry.get("institution", entry.get("school", ""))
        core = f"{authors}, {year}. {title}. {entry.get('type', 'Thesis')}, {inst}."
    elif kind == "dataset":
        core = f"{authors}, {year}. {title}. {entry.get('publisher', 'Dataset')}."
    else:
        venue = entry.get("publisher", entry.get("institution", ""))
        core = f"{authors}, {year}. {title}. {venue}."
    doi = entry.get("doi", "")
    url = entry.get("url", "")
    if doi:
        core += f" https://doi.org/{doi}."
    elif url:
        core += f" {url}."
    return re.sub(r"\s+", " ", core).strip()


def citation_order(text: str) -> list[str]:
    order: list[str] = []
    for group in re.findall(r"\[([^\]]*@[^\]]+)\]", text):
        for key in re.findall(r"@([A-Za-z0-9_]+)", group):
            if key not in order:
                order.append(key)
    return order


def replace_citations(text: str, bib: dict[str, dict[str, str]]) -> str:
    def repl(match):
        keys = re.findall(r"@([A-Za-z0-9_]+)", match.group(1))
        labels = [citation_label(bib[k]) for k in keys]
        return "(" + "; ".join(labels) + ")"
    return re.sub(r"\[([^\]]*@[^\]]+)\]", repl, text)


def normalize_text(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("--", "–")
    text = text.replace("<=", "≤").replace(">=", "≥").replace("->", "→")
    subs = {
        "H2O": "H₂O", "CO2": "CO₂", "CH4": "CH₄", "C3H8": "C₃H₈",
        "H2": "H₂", "N2": "N₂", "O2": "O₂",
    }
    for old, new in subs.items():
        text = re.sub(rf"\b{old}\b", new, text)
    text = text.replace("^(-1/2)", "⁻¹ᐟ²")
    return text


def add_rich_text(paragraph, text: str, size=10.5, bold=False, italic=False):
    # Minimal markdown emphasis support.
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        p_bold, p_italic = bold, italic
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
            p_bold = True
        elif part.startswith("*") and part.endswith("*"):
            part = part[1:-1]
            p_italic = True
        run = paragraph.add_run(normalize_text(part))
        set_run_font(run, size=size, bold=p_bold, italic=p_italic)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.line_spacing = 1.12
    pf.space_after = Pt(5.5)
    pf.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)

    for name, size, before, after in (
        ("Heading 1", 13.5, 12, 5),
        ("Heading 2", 12.0, 10, 4),
        ("Heading 3", 11.0, 8, 3),
        ("Heading 4", 10.5, 7, 2),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def extract_section(lines: list[str], heading: str, next_heading: str | None) -> str:
    start = next(i for i, line in enumerate(lines) if line.strip() == heading) + 1
    end = len(lines)
    if next_heading:
        end = next(i for i in range(start, len(lines)) if lines[i].strip() == next_heading)
    return "\n".join(lines[start:end]).strip()


def add_body(doc: Document, body_lines: list[str], bib: dict[str, dict[str, str]]):
    paragraph_buffer: list[str] = []

    def flush():
        if not paragraph_buffer:
            return
        text = " ".join(x.strip() for x in paragraph_buffer).strip()
        paragraph_buffer.clear()
        if not text:
            return
        text = replace_citations(text, bib)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.2)
        add_rich_text(p, text)

    for line in body_lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = stripped[level:].strip()
            p = doc.add_paragraph(style=f"Heading {min(level - 1, 4)}")
            add_rich_text(p, normalize_text(heading), size={2: 13.5, 3: 12, 4: 11}.get(level, 10.5), bold=True)
            continue
        if stripped.startswith("- "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, replace_citations(stripped[2:], bib))
            continue
        paragraph_buffer.append(stripped)
    flush()


def load_figure_captions(path: Path) -> dict[int, str]:
    raw = path.read_text(encoding="utf-8")
    captions = {}
    for m in re.finditer(r"\*\*Figure (\d+)\.\s*(.*?)\*\*\s*(.*?)(?=\n\n\*\*Figure|\Z)", raw, re.S):
        number = int(m.group(1))
        captions[number] = f"Figure {number}. {m.group(2)} {m.group(3).strip()}"
    return captions


def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def set_repeat_table_layout(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_table_page(doc: Document, number: int, csv_path: Path, title: str):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section(section, landscape=True)
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    add_rich_text(p, f"Table {number}. {title}", size=10.5, bold=True)

    with csv_path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if number in MAIN_TABLE_SELECTION:
        full_count = len(rows) - 1
        selected = MAIN_TABLE_SELECTION[number]
        rows = [rows[0]] + [row for row in rows[1:] if row[0] in selected]
        if len(rows) != len(selected) + 1:
            found = {row[0] for row in rows[1:]}
            raise ValueError(f"Table {number} summary selection mismatch: {sorted(selected - found)}")
        note = doc.add_paragraph()
        note.paragraph_format.keep_with_next = True
        note.paragraph_format.space_after = Pt(4)
        supplement = "S1" if number == 3 else "S2"
        add_rich_text(
            note,
            f"Main-text summary of 8/{full_count} tiers; the complete table is reproduced as Supplementary Table {supplement} and archived as CSV/JSON/Markdown.",
            size=8.2,
            italic=True,
        )
    ncols = len(rows[0])
    table = doc.add_table(rows=1, cols=ncols)
    set_repeat_table_layout(table)
    repeat_table_header(table.rows[0])

    total_width = 9.75
    width_profiles = {
        1: [1.7, 4.3, 3.75],
        2: [1.45, 0.72, 0.66, 0.66, 0.66, 0.82, 0.72, 1.05, 0.85, 2.16],
        3: [1.05, 1.35, 0.50, 0.85, 1.20, 1.35, 1.25, 2.20],
        4: [1.45, 1.6, 1.35, 1.65, 1.55, 2.15],
    }
    widths = width_profiles[number]
    assert len(widths) == ncols and abs(sum(widths) - total_width) < 0.2

    for i, text in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(cell, top=True, bottom=True)
        set_cell_margins(cell, top=20 if number in {3, 4} else 80, bottom=20 if number in {3, 4} else 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_rich_text(p, text, size=7.6 if number in {3, 4} else 7.8, bold=True)

    for r_idx, row in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, bottom=r_idx == len(rows) - 1)
            set_cell_margins(cell, top=10 if number in {3, 4} else 65, bottom=10 if number in {3, 4} else 65)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i in {1, 2, 3, 4, 5, 6} and number == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(p, text, size=7.25 if number in {3, 4} else 7.6)


def add_figure_pages(doc: Document, captions: dict[int, str]):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section(section, landscape=False)
    for number in range(1, 7):
        if number > 1:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        width = Inches(6.45)
        run.add_picture(str(FIGURES[number]), width=width)
        doc_pr = run._r.xpath(".//wp:docPr")
        if doc_pr:
            doc_pr[0].set("descr", captions[number])
        p.paragraph_format.space_after = Pt(7)
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cap.paragraph_format.keep_with_next = False
        add_rich_text(cap, captions[number], size=9)


def main():
    source_text = SOURCE.read_text(encoding="utf-8")
    source_lines = source_text.splitlines()
    title = source_lines[0].lstrip("# ").strip()
    abstract = extract_section(source_lines, "## Abstract", "## Keywords")
    keywords = extract_section(source_lines, "## Keywords", "## 1. Introduction")
    start = source_lines.index("## 1. Introduction")
    body_lines = source_lines[start:]

    order = citation_order(source_text)
    bib = parse_bib(BIB)
    missing = [key for key in order if key not in bib]
    if missing:
        raise RuntimeError(f"Missing bibliography entries: {missing}")

    doc = Document()
    configure_styles(doc)
    set_section(doc.sections[0], landscape=False)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_rich_text(p, normalize_text(title), size=16, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    add_rich_text(p, "Blinded manuscript for review", size=10, italic=True)

    p = doc.add_paragraph(style="Heading 1")
    add_rich_text(p, "Abstract", size=13.5, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    add_rich_text(p, replace_citations(abstract, bib), size=10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    add_rich_text(p, "Keywords  ", size=10.5, bold=True)
    add_rich_text(p, keywords, size=10.5)

    add_body(doc, body_lines, bib)

    # References are required in the review manuscript; source uses Pandoc keys.
    p = doc.add_paragraph(style="Heading 1")
    add_rich_text(p, "References", size=13.5, bold=True)
    reference_order = sorted(order, key=lambda key: reference_sort_key(bib[key]))
    for key in reference_order:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(3)
        add_rich_text(p, reference_text(bib[key]), size=9.5)

    for number in range(1, 5):
        add_table_page(doc, number, *TABLES[number])

    add_figure_pages(doc, load_figure_captions(CAPTIONS))

    # Avoid orphaned headings and keep document properties neutral for blind review.
    doc.core_properties.title = normalize_text(title)
    doc.core_properties.subject = "Pressure swing adsorption quantum amplitude estimation resource audit"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "pressure swing adsorption; quantum amplitude estimation; process systems engineering"

    doc.save(OUTPUT)

    # Structural shipping checks.
    xml_text = "\n".join(p.text for p in doc.paragraphs)
    assert "[@" not in xml_text
    assert "`" not in xml_text
    assert len(order) == 18
    assert reference_order == sorted(reference_order, key=lambda key: reference_sort_key(bib[key]))
    assert all(path.exists() for path in FIGURES.values())
    print(f"wrote {OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)} references={len(order)}")


if __name__ == "__main__":
    main()

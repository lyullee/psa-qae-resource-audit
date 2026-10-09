"""Audit the public toPSAil/Yang H2/CO model as independent evidence.

This artifact keeps three layers separate:

* the 1997 Yang et al. experimental performance anchor;
* the public toPSAil workbooks and the validation statements in Kim and Scott;
* the frozen product-gas recycle/event definition used in this project.

The public model has the same H2/CO chemistry and a closely related two-bed,
six-step product-tank/purge topology.  It does not expose the project's
product-gas recycle fraction as an identical control and the repository does
not distribute cycle-result MAT files.  It is therefore an independent
same-chemistry model/experiment anchor, not a same-policy selection test.
"""

from __future__ import annotations

import csv
import hashlib
import json
import subprocess
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RESULTS = ROOT / "10_TEA_LCA_Python" / "results" / "pycharm_publication_enhancement"
TOPSAIL = ROOT / "10_TEA_LCA_Python" / "external" / "toPSAil"
ARCHIVE = TOPSAIL / "4_example" / "0_archive"
SOURCE_PDF = HERE / "external_data" / "topsail_2023" / "Kim_Scott_2023_toPSAil_open_manuscript.pdf"
SOURCE_TEXT = HERE / "external_data" / "topsail_2023" / "Kim_Scott_2023_toPSAil_open_manuscript.txt"
PROTOCOL = RESULTS / "PSA_QAE_CCE_PRIMARY_CASE_PROTOCOL.json"

OUT_JSON = RESULTS / "PSA_QAE_CCE_TOPSAIL_YANG_H2_PUBLIC_MODEL_AUDIT.json"
OUT_CSV = RESULTS / "PSA_QAE_CCE_TOPSAIL_YANG_H2_PUBLIC_MODEL_AUDIT.csv"
OUT_MD = RESULTS / "PSA_QAE_CCE_TOPSAIL_YANG_H2_PUBLIC_MODEL_AUDIT.md"
OUT_PNG = RESULTS / "PSA_QAE_CCE_TOPSAIL_YANG_H2_PUBLIC_MODEL_AUDIT_FIGURE.png"
OUT_PDF = RESULTS / "PSA_QAE_CCE_TOPSAIL_YANG_H2_PUBLIC_MODEL_AUDIT_FIGURE.pdf"

CASE_NAMES = ["case_study_4.0", "case_study_4.1", "case_study_4.2", "case_study_4.3"]
INPUT_FILES = {
    "configuration": "0.1_simulation_configurations.xlsm",
    "numerics": "0.2_numerical_methods.xlsm",
    "adsorbates": "1.2_adsorbate_properties.xlsm",
    "adsorbent": "1.3_adsorbent_properties.xlsm",
    "feed": "2.1_feed_stream_properties.xlsm",
    "raffinate": "2.2_raffinate_stream_properties.xlsm",
    "adsorber": "3.1_adsorber_properties.xlsm",
    "raffinate_tank": "3.3_raffinate_tank_properties.xlsm",
    "cycle": "4.1_cycle_organization.xlsm",
}

YANG_EXPERIMENT = {
    "doi": "10.1021/ie960728h",
    "citation": "Yang, Lee and Chang, Ind. Eng. Chem. Res. 36 (1997) 2789-2798",
    "system": "two-bed six-step zeolite-5A H2/CO PSA",
    "feed_h2_fraction": 0.70,
    "feed_co_fraction": 0.30,
    "low_pressure_atm": 1.0,
    "high_pressure_atm": 11.0,
    "product_h2_purity": 0.9999,
    "h2_recovery": 0.7587,
    "studied_controls": [
        "adsorption pressure",
        "purge-to-feed ratio (P/F)",
        "adsorption/purge step time",
        "pressure-equalization step time",
    ],
}

SOURCE_REPORTED = {
    "model": "nonisothermal extended Langmuir-Freundlich equilibrium, LDF mass transfer; no-pressure-drop and Ergun variants",
    "parameter_calibration": [
        "constant-pressure valve constants matched reported feed and purge flow rates",
        "pressure-changing valve constants matched reported pressure changes over step durations",
        "heat-transfer coefficients matched the experimental product recovery",
    ],
    "spatial_cstrs": 25,
    "css_cycles": {
        "isothermal_time_driven": 13,
        "nonisothermal_time_driven": 42,
        "nonisothermal_event_driven": 36,
        "nonisothermal_ergun_time_driven": 31,
    },
    "cpu_seconds_per_cycle": {
        "flow_driven_guess_check": 16.5,
        "flow_driven_no_guess_check": 71.8,
        "pressure_driven_jacobian_pattern": 9.0,
        "pressure_driven_no_pattern": 26.0,
    },
    "ergun_high_pressure_feed_pressure_drop_order_bar": "1e-4",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_data(path: Path) -> dict:
    """Read the two-row machine-facing sheet used by toPSAil."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Data(Transposed)"]
    rows = list(sheet.iter_rows(values_only=True))
    return {
        str(key): value
        for key, value in zip(rows[0], rows[1])
        if key is not None and value is not None
    }


def git_value(*args: str) -> str | None:
    try:
        result = subprocess.run(
            ["git", "-C", str(TOPSAIL), *args],
            check=True,
            text=True,
            capture_output=True,
        )
    except (OSError, subprocess.CalledProcessError):
        return None
    return result.stdout.strip()


def normalize_steps(value: str) -> list[str]:
    return [part.split("-")[0] for part in value.split()]


def extract_case(case_name: str) -> dict:
    case_root = ARCHIVE / case_name / "1_simulation_inputs"
    raw = {
        key: workbook_data(case_root / filename)
        for key, filename in INPUT_FILES.items()
    }
    cfg = raw["configuration"]
    adsorbates = raw["adsorbates"]
    feed = raw["feed"]
    product = raw["raffinate"]
    adsorber = raw["adsorber"]
    tank = raw["raffinate_tank"]
    cycle = raw["cycle"]
    input_paths = [case_root / name for name in INPUT_FILES.values()]
    return {
        "case": case_name,
        "thermal_model": "nonisothermal" if int(cfg["bool5"]) else "isothermal",
        "pressure_drop_model": "Ergun" if int(cfg["modSp6"]) == 2 else "none",
        "event_driven": bool(int(cfg["bool4"])),
        "species": [adsorbates["sCom1"], adsorbates["sCom2"]],
        "feed_composition": {
            adsorbates["sCom1"]: float(feed["yFeC1"]),
            adsorbates["sCom2"]: float(feed["yFeC2"]),
        },
        "initial_raffinate_tank_composition": {
            adsorbates["sCom1"]: float(product["yRaC1"]),
            adsorbates["sCom2"]: float(product["yRaC2"]),
        },
        "columns": int(cycle["nCols"]),
        "spatial_cstrs": int(raw["numerics"]["nVols"]),
        "configured_cycles": int(cycle["nCycles"]),
        "steps_per_cycle": int(cycle["nSteps"]),
        "step_durations_s": [float(value) for value in str(cycle["durStep"]).split()],
        "column_1_steps": normalize_steps(str(cycle["sStepCol1"])),
        "high_pressure_bar": float(adsorber["presColHigh"]),
        "low_pressure_bar": float(adsorber["presColLow"]),
        "raffinate_tank_pressure_bar": float(tank["presRaTa"]),
        "product_tank_feeds_low_pressure_purge": "Raffinate_Tank_(T-2)" in str(cycle.get("fromProdCol1", ""))
        or float(str(cycle["valProdCol1"]).split()[-1]) > 0.0,
        "feed_valve_constants_column_1": [float(value) for value in str(cycle["valFeedCol1"]).split()],
        "product_valve_constants_column_1": [float(value) for value in str(cycle["valProdCol1"]).split()],
        "input_files": [
            {
                "relative_path": str(path.relative_to(ROOT)).replace("\\", "/"),
                "sha256": sha256(path),
                "bytes": path.stat().st_size,
            }
            for path in input_paths
        ],
    }


def make_figure(payload: dict) -> None:
    local = payload["frozen_primary_case"]
    experiment = payload["yang_experiment"]
    criteria = payload["evidence_matrix"]

    fig = plt.figure(figsize=(12.0, 6.2), constrained_layout=True)
    grid = fig.add_gridspec(1, 2, width_ratios=[1.05, 1.35])
    ax0 = fig.add_subplot(grid[0, 0])
    ax1 = fig.add_subplot(grid[0, 1])

    labels = ["Purity\n(%)", "Recovery\n(%)"]
    external_values = [100 * experiment["product_h2_purity"], 100 * experiment["h2_recovery"]]
    local_values = [100 * local["purity_threshold"], 100 * local["recovery_threshold"]]
    positions = [0, 1]
    width = 0.34
    bars1 = ax0.bar([x - width / 2 for x in positions], external_values, width, color="#0f766e", label="Yang experiment")
    bars2 = ax0.bar([x + width / 2 for x in positions], local_values, width, color="#7c3aed", label="frozen event threshold")
    ax0.set_xticks(positions, labels)
    ax0.set_ylim(70, 102)
    ax0.set_ylabel("percent")
    ax0.set_title("Experimental anchor vs frozen event")
    ax0.grid(axis="y", alpha=0.22)
    ax0.legend(loc="upper right", fontsize=9)
    for bars in (bars1, bars2):
        for bar in bars:
            ax0.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.35, f"{bar.get_height():.2f}", ha="center", va="bottom", fontsize=9)
    ax0.text(
        0.5,
        0.04,
        "Yang: purity passes, recovery fails\n=> frozen joint event = false",
        transform=ax0.transAxes,
        ha="center",
        va="bottom",
        fontsize=10,
        bbox={"boxstyle": "round,pad=0.35", "facecolor": "#fef3c7", "edgecolor": "#d97706"},
    )

    ax1.set_xlim(0, 1)
    ax1.set_ylim(-0.7, len(criteria) - 0.15)
    ax1.axis("off")
    ax1.set_title("What the public H$_2$/CO evidence closes", pad=12)
    for index, item in enumerate(reversed(criteria)):
        y = index
        passed = bool(item["satisfied"])
        color = "#166534" if passed else "#b91c1c"
        face = "#dcfce7" if passed else "#fee2e2"
        box = FancyBboxPatch((0.02, y - 0.30), 0.96, 0.60, boxstyle="round,pad=0.012", facecolor=face, edgecolor=color, linewidth=1.2)
        ax1.add_patch(box)
        mark = "✓" if passed else "×"
        ax1.text(0.055, y, mark, color=color, fontsize=17, fontweight="bold", va="center")
        ax1.text(0.12, y + 0.08, item["criterion"], fontsize=10.5, fontweight="bold", va="center")
        ax1.text(0.12, y - 0.12, item["detail"], fontsize=8.8, color="#334155", va="center")
    ax1.text(
        0.5,
        -0.58,
        "Independent same-chemistry model + experiment anchor; same-policy selection gate remains open",
        ha="center",
        va="center",
        fontsize=9.3,
        color="#0f172a",
    )

    fig.suptitle("Public toPSAil/Yang evidence strengthens process validity without validating recycle 0.420", fontsize=14, fontweight="bold")
    fig.savefig(OUT_PNG, dpi=260)
    fig.savefig(OUT_PDF)
    plt.close(fig)


def main() -> None:
    protocol = json.loads(PROTOCOL.read_text(encoding="utf-8"))
    cases = [extract_case(name) for name in CASE_NAMES]
    inner = protocol["inner_estimand"]
    local = {
        "policy_variable": protocol["outer_decision"]["variable"],
        "candidate_values": protocol["outer_decision"]["candidate_values"],
        "purity_threshold": float(inner["purity_threshold"]),
        "recovery_threshold": float(inner["recovery_threshold"]),
        "chance_constraint_target": float(protocol["chance_constraint_target"]),
    }

    topology_checks = {
        "all_cases_h2_co": all(case["species"] == ["H2", "CO"] for case in cases),
        "all_cases_feed_70_30": all(case["feed_composition"] == {"H2": 0.7, "CO": 0.3} for case in cases),
        "all_cases_two_bed_six_step": all(case["columns"] == 2 and case["steps_per_cycle"] == 6 for case in cases),
        "all_cases_product_tank_purge": all(case["product_tank_feeds_low_pressure_purge"] for case in cases),
        "all_cases_25_cstr": all(case["spatial_cstrs"] == 25 for case in cases),
    }
    experimental_joint_event = (
        YANG_EXPERIMENT["product_h2_purity"] >= local["purity_threshold"]
        and YANG_EXPERIMENT["h2_recovery"] >= local["recovery_threshold"]
    )
    mat_outputs = list(ARCHIVE.glob("case_study_4*/**/*.mat"))

    evidence_matrix = [
        {
            "criterion": "Independent H2 separation chemistry",
            "satisfied": topology_checks["all_cases_h2_co"],
            "detail": "H2/CO = 70/30 with zeolite 5A; external experiment and public model",
        },
        {
            "criterion": "Matched cyclic process structure",
            "satisfied": topology_checks["all_cases_two_bed_six_step"] and topology_checks["all_cases_product_tank_purge"],
            "detail": "2 beds, 6 steps, equalization, product tank and light-product purge",
        },
        {
            "criterion": "Experimental performance anchor",
            "satisfied": True,
            "detail": "99.99% H2 purity and 75.87% H2 recovery in Yang et al.",
        },
        {
            "criterion": "Published CSS/model calibration evidence",
            "satisfied": True,
            "detail": "valve/pressure/recovery calibration; 13/42/36/31-cycle CSS controls",
        },
        {
            "criterion": "Identical policy coordinate",
            "satisfied": False,
            "detail": "Yang P/F and product-tank purge are not the frozen product-gas recycle fraction",
        },
        {
            "criterion": "Frozen event and 0.420 selection reproduced",
            "satisfied": False,
            "detail": "Yang recovery is below 78.9%; repository distributes no result MAT files",
        },
    ]

    payload = {
        "schema": "psa-qae-cce-topsail-yang-h2-public-model-audit/v1",
        "status": "independent_same_chemistry_model_and_experimental_anchor_policy_mapping_open",
        "qpu_submission": False,
        "repository": {
            "url": git_value("remote", "get-url", "origin") or "https://github.com/taehunk333/toPSAil",
            "license": "GNU GPL v3",
            "working_head": git_value("rev-parse", "HEAD"),
            "fetched_origin_main_head": git_value("rev-parse", "origin/main"),
            "archived_case_directory": str(ARCHIVE.relative_to(ROOT)).replace("\\", "/"),
            "distributed_result_mat_file_count": len(mat_outputs),
        },
        "sources": {
            "topsail_article": {
                "doi": "10.1016/j.compchemeng.2023.108309",
                "citation": "Kim and Scott, Computers & Chemical Engineering 176 (2023) 108309",
                "local_pdf": str(SOURCE_PDF.relative_to(ROOT)).replace("\\", "/"),
                "pdf_sha256": sha256(SOURCE_PDF),
                "extracted_text_sha256": sha256(SOURCE_TEXT),
            },
            "yang_experiment": YANG_EXPERIMENT,
        },
        "yang_experiment": YANG_EXPERIMENT,
        "source_reported_model_results": SOURCE_REPORTED,
        "frozen_primary_case": local,
        "public_workbook_cases": cases,
        "topology_checks": topology_checks,
        "event_comparison": {
            "yang_purity_passes_frozen_threshold": YANG_EXPERIMENT["product_h2_purity"] >= local["purity_threshold"],
            "yang_recovery_passes_frozen_threshold": YANG_EXPERIMENT["h2_recovery"] >= local["recovery_threshold"],
            "yang_joint_event_under_frozen_thresholds": experimental_joint_event,
            "purity_margin_percentage_point": 100 * (YANG_EXPERIMENT["product_h2_purity"] - local["purity_threshold"]),
            "recovery_margin_percentage_point": 100 * (YANG_EXPERIMENT["h2_recovery"] - local["recovery_threshold"]),
        },
        "control_semantics": {
            "external_experimental_control": "purge-to-feed ratio P/F and step times",
            "public_model_control": "product-tank-to-bed purge valve/flow plus pressure and event settings",
            "primary_case_control": local["policy_variable"],
            "identical_control": False,
            "numeric_candidate_substitution_permitted": False,
            "reason": "P/F normalizes purge flow by feed flow, whereas the frozen r splits gross light product between external product and recycle; equality needs synchronized stream totals and is not implied by topology.",
        },
        "evidence_matrix": evidence_matrix,
        "gate": {
            "same_chemistry_external_experiment": True,
            "same_chemistry_open_dynamic_model": True,
            "experimental_performance_anchor": True,
            "same_control_semantics": False,
            "distributed_numeric_cycle_outputs": bool(mat_outputs),
            "same_policy_event_validation": False,
            "counts_as_external_0420_selection_validation": False,
        },
        "claim_boundary": {
            "supported": [
                "independent H2/CO two-bed six-step PSA experimental performance anchor",
                "public model structure with product tank, equalization, purge, valve calibration, and CSS evidence",
                "same-chemistry check that the frozen purity threshold is experimentally reachable",
            ],
            "unsupported": [
                "external confirmation of recycle 0.420",
                "equivalence of purge-to-feed ratio and product-gas recycle fraction",
                "external reproduction of the frozen joint event probability",
                "industrial H2 plant or historian validation",
            ],
        },
        "next_gate": "run or obtain a calibrated H2 model/experiment that varies the identical product-gas recycle split, records feed/product/purge totals, and resolves 0.420/0.427/0.430 against the frozen purity/recovery event",
    }

    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    rows = []
    for case in cases:
        rows.append(
            {
                "case": case["case"],
                "thermal_model": case["thermal_model"],
                "pressure_drop_model": case["pressure_drop_model"],
                "species": "/".join(case["species"]),
                "feed_H2_fraction": case["feed_composition"]["H2"],
                "columns": case["columns"],
                "steps": case["steps_per_cycle"],
                "spatial_CSTRs": case["spatial_cstrs"],
                "high_pressure_bar": case["high_pressure_bar"],
                "low_pressure_bar": case["low_pressure_bar"],
                "configured_cycles": case["configured_cycles"],
                "product_tank_purge": case["product_tank_feeds_low_pressure_purge"],
                "same_policy_semantics": False,
            }
        )
    with OUT_CSV.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    event = payload["event_comparison"]
    lines = [
        "# Public toPSAil/Yang H2 model and experiment audit",
        "",
        "This audit connects the public toPSAil Yang case inputs, the published model-validation statements, and the original H2/CO experiment to the frozen CCE case. It is an independent same-chemistry model/experiment anchor. It is not counted as validation of the frozen product-gas recycle selection.",
        "",
        "## Public model facts",
        "",
        "| item | audited value |",
        "|---|---|",
        "| chemistry | H2/CO = 0.70/0.30 |",
        "| adsorbent and process | zeolite 5A; two beds; six steps |",
        "| pressure range | 1.01325-11.14575 bar in the archived inputs |",
        "| cycle | equalization, repressurization, high-pressure feed/product, equalization, depressurization, low-pressure product-gas purge |",
        "| dynamic models | isothermal/nonisothermal and no-pressure-drop/Ergun variants; 25 CSTRs |",
        "| experiment | 99.99% H2 purity and 75.87% H2 recovery for H2/CO |",
        "| reported CSS controls | 13 cycles isothermal; 42 nonisothermal; 36 event-driven; 31 Ergun |",
        "| result files in repository | no distributed MAT cycle outputs found |",
        "",
        "## Frozen-event comparison",
        "",
        f"The Yang experiment exceeds the frozen purity threshold by `{event['purity_margin_percentage_point']:+.3f}` percentage point but misses the recovery threshold by `{event['recovery_margin_percentage_point']:+.2f}` percentage points. Its frozen joint-event label would therefore be **{event['yang_joint_event_under_frozen_thresholds']}**. This is a useful high-purity experimental anchor and a negative event-boundary check.",
        "",
        "## Claim boundary",
        "",
        "The workbooks independently establish the same H2/CO chemistry, two-bed six-step schedule, equalization, product tank, purge path, valve constants, and spatial discretization used in the published validation. Kim and Scott report tuning constant-pressure valves to experimental feed/purge flows, pressure-changing valves to stepwise pressure changes, and heat-transfer coefficients to experimental recovery.",
        "",
        "The external P/F coordinate is not the local product-gas recycle fraction. The former is purge flow divided by feed flow; the latter is the fraction of gross light product split from external product and returned internally. The numeric candidates `0.420/0.427/0.430` cannot be assigned to P/F without synchronized stream totals. Consequently, this closes a same-chemistry model/experiment gap but leaves the same-policy selection gate open.",
        "",
        f"Status: `{payload['status']}`.",
    ]
    OUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    make_figure(payload)

    print(
        json.dumps(
            {
                "status": payload["status"],
                "cases": len(cases),
                "topology_checks_pass": all(topology_checks.values()),
                "yang_joint_event": experimental_joint_event,
                "same_policy_validation": payload["gate"]["same_policy_event_validation"],
                "qpu_submission": False,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
